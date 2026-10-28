"""架空の会社の元帳・通帳と、JEVの代わりの模擬サーバーで全体の動きを確かめる。

実行: python skills/bank-journal/tests/test_sample.py
"""

import csv
import datetime as dt
import json
import os
import subprocess
import sys
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))

from openpyxl import load_workbook  # noqa: E402

LEDGER_HEADER = ["日付", "相手勘定科目", "相手補助科目", "税区分", "摘要", "入金額", "出金額"]


def make_ledger(path: Path) -> None:
    rows = []
    balance, rate = 10_000_000, 0.015
    prev = dt.date(2025, 3, 10)
    for m in range(12):
        y, mo = (2025, 4 + m) if m < 9 else (2026, m - 8)
        d = lambda day: dt.date(y, mo, day).isoformat()  # noqa: E731
        rows.append([d(5), "売掛金", "青山工業", "対象外", "ｱｵﾔﾏｺｳｷﾞﾖｳ(ｶ", 300_000 + m * 10_000, ""])
        rows.append([d(26), "通信費", "", "課対仕入10%", f"ﾄﾞｺﾓ ﾂｳｼﾝﾘﾖｳ {mo}月分", "", 9_800 + m * 50])
        rows.append([d(27), "地代家賃", "", "課対仕入10%", "ﾌﾘｺﾐ ﾔﾏﾀﾞﾌﾄﾞｳｻﾝ", "", 165_000])
        rows.append([d(27), "支払手数料", "", "課対仕入10%", "振込手数料", "", 660])
        rows.append([d(25), "未払金", "給与", "対象外", "給与 総合振込", "", 1_200_000 + m * 1_000])
        rows.append([d(25), "支払手数料", "", "課対仕入10%", "振込手数料", "", 990])
        rows.append([d(28), "買掛金", "丸々商事", "対象外", "丸々商事 支払", "", 200_000 + m * 7_000])
        # 元利均等返済（毎月105,000円、利息は実日数で計算）
        pay = dt.date(y, mo, 10)
        interest = round(balance * rate * (pay - prev).days / 365)
        principal = 105_000 - interest
        balance -= principal
        prev = pay
        rows.append([pay.isoformat(), "長期借入金", "〇〇銀行", "対象外", "ﾕｳｼﾍﾝｻｲ", "", principal])
        rows.append([pay.isoformat(), "支払利息", "", "非課税", "ﾕｳｼﾍﾝｻｲ", "", interest])
    rows.append(["2025-06-15", "現金", "", "対象外", "ｶｰﾄﾞ ATM", "", 50_000])
    rows.append(["2025-09-15", "役員借入金", "", "対象外", "ｶｰﾄﾞ ATM", "", 30_000])
    write_csv(path, LEDGER_HEADER, rows)
    return balance, prev


def make_passbook(path: Path, loan_balance: int, loan_prev: dt.date) -> int:
    expected_interest = round(loan_balance * 0.015 * (dt.date(2026, 4, 10) - loan_prev).days / 365)
    rows = [
        ["2026-04-01", "繰越", "", "", "5,000,000", ""],
        ["2026-04-06", "ｱｵﾔﾏｺｳｷﾞﾖｳ(ｶ", "", "420,000", "5,420,000", ""],
        ["2026-04-10", "ﾕｳｼﾍﾝｻｲ", "105,000", "", "5,315,000", ""],
        ["2026-04-15", "ｶｰﾄﾞ ATM", "40,000", "", "5,275,000", ""],
        ["2026-04-20", "ﾌﾘｺﾐ ｻｸﾗﾃﾞﾝｷ(ｶ", "88,000", "", "5,187,000", ""],
        ["2026-04-25", "ｷﾕｳﾖ ｿｳｺﾞｳﾌﾘｺﾐ", "1,213,990", "", "3,973,010", ""],
        ["2026-04-27", "ﾄﾞｺﾓ ﾂｳｼﾝﾘﾖｳ", "10,450", "", "3,962,560", ""],
        ["2026-04-27", "ﾌﾘｺﾐ ﾔﾏﾀﾞﾌﾄﾞｳｻﾝ", "165,000", "", "3,797,560", ""],
        ["2026-04-27", "ﾌﾘｺﾐﾃｽｳﾘﾖｳ", "660", "", "3,796,900", ""],
        ["2026-04-28", "ﾌﾘｺﾐ ﾏﾙﾏﾙｼﾖｳｼﾞ(ｶ", "284,000", "", "3,512,900", ""],
        ["2026-04-30", "ﾘｿｸ", "", "35", "3,519,935", "残高の末尾が読みにくい"],  # わざと残高を誤読
    ]
    write_csv(path, ["日付", "摘要", "出金額", "入金額", "残高", "読取メモ"], rows)
    return expected_interest


def write_csv(path, header, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


class MockJev(BaseHTTPRequestHandler):
    """カナの相手先名と科目の対応表で答える、JEVの代わり。"""

    answers = {"ｷﾕｳﾖ": "未払金／給与", "ﾏﾙﾏﾙｼﾖｳｼﾞ": "買掛金／丸々商事"}
    requests = []

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        assert self.headers["Authorization"] == "Bearer test-key"
        MockJev.requests.append(body)
        q = body["questions"]["answer"]
        text = body["state"]["通帳の取引"]["摘要"]
        choice = "UNKNOWN"
        for word, account in self.answers.items():
            if word in text:
                choice = next((k for k, v in q["criteria"].items() if v.startswith(account)), "UNKNOWN")
        probs = {k: (0.92 if k == choice else 0.08 / (len(q["criteria"]) - 1)) for k in q["criteria"]}
        data = json.dumps({"model": "jev-mock", "usage": {},
                           "answers": {"answer": {"type": "choice", "choice": choice,
                                                  "confidence": 0.92, "probabilities": probs}}}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, *args):
        pass


def run(*args, env=None):
    result = subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True, env=env, cwd=SCRIPTS)
    print(result.stdout + result.stderr)
    assert result.returncode == 0, result.stderr
    return result.stdout


def read_journal(path):
    ws = load_workbook(path)["仕訳"]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]


def main():
    tmp = Path(tempfile.mkdtemp())
    loan_balance, loan_prev = make_ledger(tmp / "ledger.csv")
    expected_interest = make_passbook(tmp / "passbook.csv", loan_balance, loan_prev)
    run(SCRIPTS / "build_rules.py", tmp / "ledger.csv", tmp / "rules.json")

    # 1) JEVなし：分からないものは「不明」
    run(SCRIPTS / "journalize.py", tmp / "passbook.csv", tmp / "rules.json", tmp / "nojev.xlsx",
        "--bank-sub", "〇〇銀行", "--no-jev")
    rows = read_journal(tmp / "nojev.xlsx")
    by_text = {}
    for r in rows:
        by_text.setdefault(r["摘要"], []).append(r)
    assert by_text["ｱｵﾔﾏｺｳｷﾞﾖｳ(ｶ"][0]["貸方科目"] == "売掛金"
    assert by_text["ｱｵﾔﾏｺｳｷﾞﾖｳ(ｶ"][0]["判定"] == "過去一致"
    assert by_text["ﾄﾞｺﾓ ﾂｳｼﾝﾘﾖｳ"][0]["借方科目"] == "通信費"
    assert by_text["ﾌﾘｺﾐﾃｽｳﾘﾖｳ"][0]["借方科目"] == "支払手数料"
    assert by_text["ﾌﾘｺﾐ ｻｸﾗﾃﾞﾝｷ(ｶ"][0]["判定"] == "不明"
    assert by_text["ｶｰﾄﾞ ATM"][0]["判定"] == "不明", "現金と役員借入金の両方がある摘要は不明にする"
    assert by_text["ﾘｿｸ"][0]["判定"] == "不明"
    # 家賃は通帳に手数料の行があるので分割しない
    assert [r["金額"] for r in by_text["ﾌﾘｺﾐ ﾔﾏﾀﾞﾌﾄﾞｳｻﾝ"]] == [165_000]
    # 借入返済：元本と利息に分割
    loan = by_text["ﾕｳｼﾍﾝｻｲ"]
    assert [r["借方科目"] for r in loan] == ["長期借入金", "支払利息"]
    assert sum(r["金額"] for r in loan) == 105_000
    assert abs(loan[1]["金額"] - expected_interest) <= 3, (loan[1]["金額"], expected_interest)
    assert loan[0]["判定"] == "推定"
    # 残高の誤読を検出
    assert rows[-1]["残高チェック"].startswith("不一致")

    # 2) 模擬JEVあり：カナと漢字の違う相手先もJEVが判定
    server = HTTPServer(("127.0.0.1", 0), MockJev)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    env = {**os.environ, "TYPESAFE_API_KEY": "test-key", "TYPESAFE_BASE_URL": f"http://127.0.0.1:{server.server_port}"}
    run(SCRIPTS / "journalize.py", tmp / "passbook.csv", tmp / "rules.json", tmp / "jev.xlsx",
        "--bank-sub", "〇〇銀行", env=env)
    rows = read_journal(tmp / "jev.xlsx")
    by_text = {}
    for r in rows:
        by_text.setdefault(r["摘要"], []).append(r)
    maru = by_text["ﾌﾘｺﾐ ﾏﾙﾏﾙｼﾖｳｼﾞ(ｶ"]
    assert maru[0]["借方科目"] == "買掛金" and maru[0]["判定"] == "JEV判定", maru
    salary = by_text["ｷﾕｳﾖ ｿｳｺﾞｳﾌﾘｺﾐ"]
    assert [(r["借方科目"], r["金額"]) for r in salary] == [("未払金", 1_213_000), ("支払手数料", 990)], salary
    assert salary[0]["判定"] == "推定"
    assert by_text["ﾌﾘｺﾐ ｻｸﾗﾃﾞﾝｷ(ｶ"][0]["判定"] == "不明"
    assert by_text["ｱｵﾔﾏｺｳｷﾞﾖｳ(ｶ"][0]["判定"] == "過去一致", "過去一致はJEVに聞かない"
    asked = {r["state"]["通帳の取引"]["摘要"] for r in MockJev.requests}
    assert "ｱｵﾔﾏｺｳｷﾞﾖｳ(ｶ" not in asked

    # 3) 通帳キーワードの追加で、漢字の元帳でも過去一致にできる
    write_csv(tmp / "keywords.csv", ["相手勘定科目", "相手補助科目", "通帳キーワード"], [["買掛金", "丸々商事", "ﾏﾙﾏﾙｼﾖｳｼﾞ"]])
    run(SCRIPTS / "build_rules.py", tmp / "ledger.csv", tmp / "rules2.json", "--keywords", tmp / "keywords.csv")
    run(SCRIPTS / "journalize.py", tmp / "passbook.csv", tmp / "rules2.json", tmp / "kw.xlsx", "--no-jev")
    maru = [r for r in read_journal(tmp / "kw.xlsx") if r["摘要"] == "ﾌﾘｺﾐ ﾏﾙﾏﾙｼﾖｳｼﾞ(ｶ"]
    assert maru[0]["借方科目"] == "買掛金" and maru[0]["判定"] == "過去一致"

    server.shutdown()
    print("すべてのテストに合格しました。出力例:", tmp / "jev.xlsx")


if __name__ == "__main__":
    main()
