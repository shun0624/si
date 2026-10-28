"""コマ表（manga/コマ表.xlsx）を作る。

python manga/scripts/make_sheet.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "コマ表.xlsx"
ROWS = 300  # 式を入れておく行数

HEAD = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")
AUTO = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, titles, widths):
    for col, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=1, column=col, value=t)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, WRAP
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def style(ws, rows, cols, fill):
    for r in range(2, rows + 2):
        for c in cols:
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.border, cell.alignment = fill, BOX, WRAP


def build():
    wb = Workbook()

    # 使い方
    ws = wb.active
    ws.title = "使い方"
    lines = [
        "コマ表の使い方",
        "",
        "1. 「キャラ」シートに、キャラの名前・呼び出し名・見た目の特徴・LoRAのファイル名を書く",
        "2. 「コマ表」シートに、ページ・コマ番号・出るキャラ・表情・場所などを書く（黄色い列）",
        "3. 緑の列（指示文・保存名・使うLoRA）が自動でできる",
        "4. ComfyUIで「指示文」を「指示文」欄に、「保存名」を「保存名」欄に貼る",
        "5. 「使うLoRA」に出ているLoRAをオンにして実行する",
        "6. 選んだ画像のファイル名を「採用した画像」に書いておく",
        "",
        "・ページは、ネーム画像のファイル名と同じにする（p1.png なら p1）",
        "・コマ番号は、コマ切り出しツールの確認用画像の番号と同じにする",
        "・表情や場所は英語の短い単語をカンマで区切って書く（例: smile, looking at viewer）",
        "・「共通設定」シートの画質タグは全部のコマに付く",
        "・キャラの特徴には人数（1girl など）を書かない（人数はコマ表の「人数タグ」に書く）",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 100

    # 共通設定
    cs = wb.create_sheet("共通設定")
    header(cs, ["項目", "内容"], [18, 100])
    cs["A2"], cs["B2"] = "画質タグ", "masterpiece, best quality, amazing quality, monochrome, greyscale, manga, lineart, screentone"
    cs["A3"], cs["B3"] = "描かせないもの", "bad quality, worst quality, worst detail, lowres, bad anatomy, bad hands, censor, signature, watermark, username, text, speech bubble, frame, border, color"
    cs["A4"], cs["B4"] = "保存フォルダ", "manga"
    style(cs, 3, [2], INPUT)

    # キャラ
    ch = wb.create_sheet("キャラ")
    header(ch, ["キャラ名", "呼び出し名（LoRAのトリガー）", "見た目の特徴（英語タグ）",
                "LoRAファイル名", "LoRAの強さ"], [14, 22, 60, 30, 12])
    examples = [
        ("あかね", "akane01", "short black hair, bob cut, red eyes, hair clip, school uniform, red ribbon", "akane01.safetensors", 0.9),
        ("そうた", "souta01", "messy brown hair, glasses, gakuran", "souta01.safetensors", 0.9),
    ]
    for r, row in enumerate(examples, 2):
        for c, v in enumerate(row, 1):
            ch.cell(row=r, column=c, value=v)
    style(ch, 30, range(1, 6), INPUT)

    # コマ表
    ks = wb.create_sheet("コマ表")
    titles = ["ページ", "コマ", "キャラ1", "キャラ2", "人数タグ", "表情・動き", "場所・背景",
              "カメラ・構図", "追加の指示", "指示文（自動）", "保存名（自動）", "使うLoRA（自動）",
              "採用した画像", "メモ"]
    widths = [8, 6, 12, 12, 14, 30, 24, 22, 22, 70, 18, 30, 26, 24]
    header(ks, titles, widths)
    samples = [
        ("p1", 1, "あかね", "", "1girl, solo", "surprised, looking back", "classroom", "upper body", ""),
        ("p1", 2, "あかね", "そうた", "1girl, 1boy", "talking, smile", "school hallway", "from side", ""),
    ]
    for r, row in enumerate(samples, 2):
        for c, v in enumerate(row, 1):
            if v != "":
                ks.cell(row=r, column=c, value=v)

    def look(col, n):
        # キャラ名からキャラシートの値を引く（見つからなければ空）
        return f'IFERROR(VLOOKUP({col}{{r}},キャラ!$A:$E,{n},FALSE),"")'

    for r in range(2, ROWS + 2):
        parts = ["共通設定!$B$2", f"E{r}", look("C", 2), look("C", 3), look("D", 2), look("D", 3),
                 f"F{r}", f"G{r}", f"H{r}", f"I{r}"]
        parts = [p.replace("{r}", str(r)) for p in parts]
        ks[f"J{r}"] = (f'=IF(AND(C{r}="",F{r}=""),"",'
                       f'_xlfn.TEXTJOIN(", ",TRUE,{",".join(parts)}))')
        ks[f"K{r}"] = (f'=IF(OR(A{r}="",B{r}=""),"",'
                       f'共通設定!$B$4&"/"&A{r}&"_c"&TEXT(B{r},"00"))')
        a = look("C", 4).replace("{r}", str(r))
        b = look("D", 4).replace("{r}", str(r))
        ka = look("C", 5).replace("{r}", str(r))
        # 2人のコマは強さを 0.5 にして、あとで1人ずつ描き直す
        ks[f"L{r}"] = (f'=IF({a}="","",{a}&"（"&IF(D{r}="",{ka},0.5)&"）")'
                       f'&IF({b}="",""," / "&{b}&"（0.5）")')
    style(ks, ROWS, [1, 2, 3, 4, 5, 6, 7, 8, 9, 13, 14], INPUT)
    style(ks, ROWS, [10, 11, 12], AUTO)

    names = DataValidation(type="list", formula1="=キャラ!$A$2:$A$31", allow_blank=True)
    ks.add_data_validation(names)
    names.add(f"C2:D{ROWS + 1}")

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
