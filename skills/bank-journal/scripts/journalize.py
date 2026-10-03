"""通帳の取引を、ルール表（build_rules.py の出力）とJEVで仕訳にしてExcelに出す。

使い方:
  python journalize.py passbook.csv rules.json 出力.xlsx --bank-account 普通預金 --bank-sub ○○銀行

passbook.csv の列: 日付, 摘要, 出金額, 入金額, 残高, 読取メモ
判定の種類:
  過去一致 … 過去の元帳に同じ相手先があり、科目が1つに決まる
  JEV判定 … JEVが候補から選び、確信度がしきい値以上
  推定     … 科目は分かるが、金額の内訳（元本と利息、手数料）を推定した
  要確認   … 候補はあるが自信が低い。必ず確認する
  不明     … 判断できない。科目は「不明」のまま出す
"""

import argparse
import datetime as dt
import difflib
import json
import sys
from concurrent.futures import ThreadPoolExecutor

from common import PASSBOOK_COLUMNS, is_fee_text, match_key, parse_amount, parse_date, read_csv
from jev_client import JevClient, JevUnavailable

FEE_MAX = 1100
FEE_ACCOUNT = "支払手数料"
UNKNOWN = "不明"
MAX_CANDIDATES = 12
STATUS_ORDER = ["過去一致", "JEV判定", "推定", "要確認", UNKNOWN]


# ---------------------------------------------------------------- 通帳の読み込みと残高チェック

def load_passbook(path: str) -> list[dict]:
    rows = []
    running = None
    for i, r in enumerate(read_csv(path, PASSBOOK_COLUMNS[:5]), start=1):
        try:
            withdrawal, deposit = parse_amount(r["出金額"]), parse_amount(r["入金額"])
            balance = parse_amount(r["残高"]) if (r.get("残高") or "").strip() else None
            date = parse_date(r["日付"])
        except ValueError as e:
            raise SystemExit(f"通帳 {i}行目: {e}")
        check = ""
        if running is not None:
            running += deposit - withdrawal
        if balance is not None:
            if running is not None and running != balance:
                check = f"不一致（計算 {running:,} / 通帳 {balance:,}、差 {balance - running:+,}）"
            elif running is not None:
                check = "OK"
            running = balance
        rows.append({
            "no": i, "date": date, "text": (r["摘要"] or "").strip(),
            "withdrawal": withdrawal, "deposit": deposit, "balance": balance,
            "memo": (r.get("読取メモ") or "").strip(), "check": check,
        })
    return rows


# ---------------------------------------------------------------- 候補の絞り込み

def name_similarity(row_key: str, pattern: dict) -> float:
    best = 0.0
    if not row_key:
        return best
    for key in pattern["keys"]:
        if key == row_key:
            return 1.0
        short, long_ = sorted((key, row_key), key=len)
        if len(short) >= 3 and short in long_:
            best = max(best, 0.9)
        else:
            best = max(best, difflib.SequenceMatcher(None, row_key, key).ratio() * 0.85)
    return best


def amount_fit(amount: int, p: dict) -> float:
    if p["amount_fixed"]:
        return 1.0 if amount == p["amount_min"] else 0.0
    if p["amount_min"] <= amount <= p["amount_max"]:
        return 0.7
    if p["amount_median"] and 0.5 <= amount / p["amount_median"] <= 2:
        return 0.4
    return 0.0


def day_fit(day: int, p: dict) -> float:
    return 1.0 if any(abs(day - d) <= 2 for d in p["days"]) else 0.0


def rank_candidates(row: dict, patterns: list[dict]) -> list[tuple[float, float, dict]]:
    direction = "出金" if row["withdrawal"] else "入金"
    amount = row["withdrawal"] or row["deposit"]
    key = match_key(row["text"])
    fee_row = is_fee_text(row["text"]) and direction == "出金" and amount <= FEE_MAX
    ranked = []
    for p in patterns:
        if p["direction"] != direction:
            continue
        sim = name_similarity(key, p)
        if fee_row and p["account"] == FEE_ACCOUNT:
            sim = 1.0
        score = 0.7 * sim + 0.2 * amount_fit(amount, p) + 0.1 * day_fit(row["date"].day, p)
        ranked.append((score, sim, p))
    ranked.sort(key=lambda t: -t[0])
    return ranked


def describe(p: dict) -> str:
    name = p["account"] + (f"／{p['sub']}" if p["sub"] else "")
    amount = (f"毎回{p['amount_min']:,}円" if p["amount_fixed"]
              else f"{p['amount_min']:,}〜{p['amount_max']:,}円（中央値{p['amount_median']:,}円）")
    return (f"{name}（税区分: {p['tax'] or 'なし'}）。過去{p['count']}回。"
            f"過去の摘要例: {' / '.join(p['examples'][:3])}。金額: {amount}。"
            f"主な日付: {', '.join(str(d) for d in p['days'][:6])}日")


# ---------------------------------------------------------------- 判定

def decide(row: dict, patterns: list[dict], jev: "JevClient | None", threshold: float, bank_label: str) -> dict:
    ranked = rank_candidates(row, patterns)
    if not ranked:
        return {"status": UNKNOWN, "pattern": None, "confidence": None, "reason": "過去に同じ入出金区分の仕訳がない"}

    exact = [(s, p) for s, sim, p in ranked if sim >= 0.9]
    # 借入返済と同じ摘要の支払利息は、返済の一部として扱う（元本と利息は後で分ける）
    linked = {p["loan"]["interest_pattern_id"] for _, p in exact if p.get("loan")}
    exact = [(s, p) for s, p in exact if p["id"] not in linked]
    if len(exact) == 1:
        p = exact[0][1]
        return {"status": "過去一致", "pattern": p, "confidence": 1.0, "reason": f"過去の摘要と一致（{p['examples'][0]}）"}
    if len(exact) > 1:
        amount = row["withdrawal"] or row["deposit"]
        fitting = [p for _, p in exact if amount_fit(amount, p) >= 0.7]
        if len(fitting) == 1:
            p = fitting[0]
            return {"status": "過去一致", "pattern": p, "confidence": 0.95,
                    "reason": f"過去の摘要と一致し、金額の範囲から {p['account']} に特定"}
        candidates = [p for _, p in exact]
    else:
        candidates = [p for _, _, p in ranked[:MAX_CANDIDATES]]

    if jev is not None:
        return ask_jev(row, candidates, jev, threshold, bank_label)
    return local_fallback(ranked, exact)


def ask_jev(row: dict, candidates: list[dict], jev: JevClient, threshold: float, bank_label: str) -> dict:
    criteria = {p["id"]: describe(p) for p in candidates}
    criteria["UNKNOWN"] = "どの候補にも当てはまらない、または情報が足りず判断できない"
    state = {
        "銀行口座": bank_label,
        "通帳の取引": {
            "日付": row["date"].isoformat(),
            "摘要": row["text"],
            "区分": "出金" if row["withdrawal"] else "入金",
            "金額": row["withdrawal"] or row["deposit"],
        },
    }
    answer = jev.choice(
        state,
        "この通帳の取引は、過去の仕訳パターンのうちどれと同じ相手先・同じ種類の取引か。"
        "摘要の相手先名（カナ・漢字・略称の違いを含む）と金額・日付の傾向で判断する。",
        criteria,
    )
    conf = float(answer["confidence"])
    if answer["choice"] == "UNKNOWN":
        return {"status": UNKNOWN, "pattern": None, "confidence": conf, "reason": "JEV: 該当する過去パターンなし"}
    p = next(c for c in candidates if c["id"] == answer["choice"])
    runner_up = sorted(((v, k) for k, v in answer["probabilities"].items() if k != answer["choice"]), reverse=True)
    alt = ""
    if runner_up and runner_up[0][1] != "UNKNOWN":
        q = next(c for c in candidates if c["id"] == runner_up[0][1])
        alt = f"／次点 {q['account']}{'／' + q['sub'] if q['sub'] else ''}（{runner_up[0][0]:.0%}）"
    if conf >= threshold:
        return {"status": "JEV判定", "pattern": p, "confidence": conf, "reason": f"JEV 確信度{conf:.0%}{alt}"}
    if conf >= 0.5:
        return {"status": "要確認", "pattern": p, "confidence": conf, "reason": f"JEV 確信度が低い（{conf:.0%}）{alt}"}
    return {"status": UNKNOWN, "pattern": None, "confidence": conf,
            "reason": f"JEV でも判断できず（最有力 {p['account']} {conf:.0%}）"}


def local_fallback(ranked, exact) -> dict:
    """JEVを使わない場合。似ている候補を示すだけで、確定はしない。"""
    if exact:
        names = "、".join(f"{p['account']}{'／' + p['sub'] if p['sub'] else ''}" for _, p in exact[:3])
        return {"status": UNKNOWN, "pattern": None, "confidence": None, "reason": f"同じ摘要で複数の科目あり（{names}）"}
    score, sim, p = ranked[0]
    second = ranked[1][0] if len(ranked) > 1 else 0
    if sim >= 0.6 and score >= 0.6 and score - second >= 0.15:
        return {"status": "要確認", "pattern": p, "confidence": round(score, 2),
                "reason": f"摘要が似ている（{p['examples'][0]}）"}
    return {"status": UNKNOWN, "pattern": None, "confidence": None, "reason": "過去に似た摘要がない"}


# ---------------------------------------------------------------- 内訳の推定（借入返済・振込手数料）

def estimate_loan(loan: dict, date: dt.date, amount: int) -> tuple[int, int, str, bool]:
    """戻り値: (元本, 利息, 説明, 自信あり)。loan["history"] に推定結果を追記する。"""
    hist = loan["history"]
    if not hist:
        return amount, 0, "過去の利息の記録がないため全額を元本にした", False
    last = hist[-1]
    last_date = dt.date.fromisoformat(last["date"])
    days = (date - last_date).days if date > last_date else 30

    def interest_by_days() -> int:
        if len(hist) >= 2:
            prev_days = (last_date - dt.date.fromisoformat(hist[-2]["date"])).days or 30
            daily = last["interest"] / prev_days
            if len(hist) >= 3 and hist[-2]["interest"]:
                prev2_days = (dt.date.fromisoformat(hist[-2]["date"]) - dt.date.fromisoformat(hist[-3]["date"])).days or 30
                decline = daily / (hist[-2]["interest"] / prev2_days)
                daily *= min(1.0, max(0.9, decline))
            return round(daily * days)
        return last["interest"]

    recent = hist[-3:]
    principals = [h["principal"] for h in recent]
    totals = [h["principal"] + h["interest"] for h in recent]
    ok = len(recent) >= 2
    if ok and max(principals) - min(principals) <= 1:
        principal, interest = principals[-1], amount - principals[-1]
        note = "元金均等返済（元本は前回と同額、利息＝引落額－元本）"
        ok = 0 <= interest <= last["interest"] * 1.5 + 100
        if not ok:
            note += "。利息が前回と大きく違うため要確認"
    elif ok and max(totals) - min(totals) <= 1:
        interest = interest_by_days()
        principal = amount - interest
        note = f"元利均等返済（利息を前回の日割り×{days}日で推定）"
        if abs(amount - totals[-1]) > 1:
            ok, note = False, note + "。引落額が前回と違うため要確認（金利・条件の変更？）"
    else:
        interest = interest_by_days()
        principal = amount - interest
        ok, note = False, "返済方式を特定できないため、利息を日割りで仮に推定"
    if principal < 0 or interest < 0:
        return amount, 0, "推定がマイナスになったため全額を元本にした。返済予定表を確認", False
    hist.append({"date": date.isoformat(), "principal": principal, "interest": interest})
    return principal, interest, note, ok


# ---------------------------------------------------------------- 仕訳の組み立て

def make_lines(row, result, bank_account, bank_sub, fee_pattern, fee_dates) -> list[dict]:
    amount = row["withdrawal"] or row["deposit"]
    outgoing = bool(row["withdrawal"])
    p = result["pattern"]
    base = {
        "no": row["no"], "date": row["date"], "text": row["text"], "check": row["check"],
        "confidence": result["confidence"],
    }

    def line(account, sub, tax, value, status, reason):
        counter = {"科目": account, "補助": sub, "税区分": tax}
        bank = {"科目": bank_account, "補助": bank_sub, "税区分": ""}
        debit, credit = (counter, bank) if outgoing else (bank, counter)
        return {**base, "debit": debit, "credit": credit, "amount": value, "status": status, "reason": reason}

    if p is None:
        return [line(UNKNOWN, "", "", amount, UNKNOWN, result["reason"])]

    status, reason = result["status"], result["reason"]
    if p.get("loan") is not None and outgoing:
        principal, interest, note, ok = estimate_loan(p["loan"], row["date"], amount)
        split_status = status if status in ("要確認",) else ("推定" if ok else "要確認")
        lines = [line(p["account"], p["sub"], p["tax"], principal, split_status, f"{reason}。{note}")]
        if interest:
            loan = p["loan"]
            lines.append(line(loan["interest_account"], loan["interest_sub"], loan["interest_tax"], interest,
                              split_status, f"借入返済の利息部分。{note}"))
        return lines

    if (outgoing and p["fee_rate"] >= 0.5 and p["fee_amount"] and amount > p["fee_amount"]
            and row["date"] not in fee_dates):
        fee = p["fee_amount"]
        fee_acct = fee_pattern or {"account": FEE_ACCOUNT, "sub": "", "tax": ""}
        split_status = "要確認" if status == "要確認" else "推定"
        note = (f"過去は支払の{p['fee_rate']:.0%}で振込手数料あり。通帳に手数料の行がないため、"
                f"引落額に手数料{fee:,}円が含まれると推定")
        return [
            line(p["account"], p["sub"], p["tax"], amount - fee, split_status, f"{reason}。{note}"),
            line(fee_acct["account"], fee_acct["sub"], fee_acct["tax"], fee, split_status, note),
        ]

    if outgoing and 0 < p["fee_rate"] < 0.5:
        reason += f"。過去に手数料込みの支払が{p['fee_rate']:.0%}あり"
    return [line(p["account"], p["sub"], p["tax"], amount, status, reason)]


# ---------------------------------------------------------------- Excel出力

def write_excel(path, lines, passbook, rules, summary):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {
        "過去一致": None,
        "JEV判定": PatternFill("solid", fgColor="E2EFDA"),
        "推定": PatternFill("solid", fgColor="FFF2CC"),
        "要確認": PatternFill("solid", fgColor="FCE4D6"),
        UNKNOWN: PatternFill("solid", fgColor="F8CBAD"),
    }
    bold = Font(bold=True)
    wb = Workbook()

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font = bold
            c.fill = PatternFill("solid", fgColor="D9E1F2")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "仕訳"
    sheet(ws, ["通帳行", "日付", "借方科目", "借方補助", "借方税区分", "貸方科目", "貸方補助", "貸方税区分",
               "金額", "摘要", "判定", "確信度", "根拠", "残高チェック"],
          [[l["no"], l["date"], l["debit"]["科目"], l["debit"]["補助"], l["debit"]["税区分"],
            l["credit"]["科目"], l["credit"]["補助"], l["credit"]["税区分"], l["amount"], l["text"], l["status"],
            None if l["confidence"] is None else round(l["confidence"], 2), l["reason"], l["check"]] for l in lines],
          [6, 11, 14, 14, 12, 14, 14, 12, 12, 28, 9, 7, 60, 30])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "yyyy/mm/dd"
        row[8].number_format = "#,##0"
        row[12].alignment = Alignment(wrap_text=True, vertical="top")
        fill = fills.get(row[10].value)
        if fill:
            for c in row:
                c.fill = fill
        if row[13].value and row[13].value != "OK":
            row[13].font = Font(color="C00000", bold=True)

    ws = wb.create_sheet("集計")
    for item in summary:
        ws.append(item)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70
    for row in ws.iter_rows():
        if row[0].value and str(row[0].value).startswith("■"):
            row[0].font = bold

    ws = wb.create_sheet("通帳読取")
    sheet(ws, ["行", "日付", "摘要", "出金額", "入金額", "残高", "読取メモ", "残高チェック"],
          [[r["no"], r["date"], r["text"], r["withdrawal"] or None, r["deposit"] or None, r["balance"], r["memo"], r["check"]]
           for r in passbook], [5, 11, 30, 12, 12, 14, 24, 40])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "yyyy/mm/dd"
        for c in row[3:6]:
            c.number_format = "#,##0"

    ws = wb.create_sheet("ルール表")
    sheet(ws, ["ID", "入出金", "相手科目", "相手補助", "税区分", "回数", "通帳照合キー", "過去の摘要例", "金額下限",
               "金額上限", "毎回同額", "主な日", "手数料込み率", "手数料額", "借入返済の履歴数"],
          [[p["id"], p["direction"], p["account"], p["sub"], p["tax"], p["count"], " / ".join(p["keys"][:6]),
            " / ".join(p["examples"][:3]), p["amount_min"], p["amount_max"], "○" if p["amount_fixed"] else "",
            ", ".join(map(str, p["days"][:8])), p["fee_rate"] or None, p["fee_amount"] or None,
            len(p["loan"]["history"]) if p.get("loan") else None] for p in rules["patterns"]],
          [6, 6, 16, 16, 12, 5, 30, 40, 11, 11, 6, 14, 8, 8, 8])
    wb.save(path)


# ---------------------------------------------------------------- メイン

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("passbook")
    ap.add_argument("rules")
    ap.add_argument("output")
    ap.add_argument("--bank-account", default="普通預金")
    ap.add_argument("--bank-sub", default="")
    ap.add_argument("--threshold", type=float, default=0.8, help="JEV判定とする確信度（既定 0.8）")
    ap.add_argument("--no-jev", action="store_true", help="JEVを使わない")
    ap.add_argument("--workers", type=int, default=8)
    args = ap.parse_args()

    with open(args.rules, encoding="utf-8") as f:
        rules = json.load(f)
    patterns = rules["patterns"]
    passbook = load_passbook(args.passbook)
    targets = [r for r in passbook if r["withdrawal"] or r["deposit"]]
    bank_label = args.bank_account + (f"／{args.bank_sub}" if args.bank_sub else "")

    jev = None if args.no_jev else JevClient.from_environment()
    jev_note = "使用しない（--no-jev）" if args.no_jev else ("使用" if jev else "APIキー未設定のため使用せず")

    def run(row):
        try:
            return decide(row, patterns, jev, args.threshold, bank_label), None
        except JevUnavailable as e:
            return decide(row, patterns, None, args.threshold, bank_label), str(e)

    started = dt.datetime.now()
    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        outcomes = list(pool.map(run, targets))
    elapsed = (dt.datetime.now() - started).total_seconds()
    errors = sorted({e for _, e in outcomes if e})
    if errors:
        jev_note = f"一部または全部で使用できず（{errors[0]}）"

    fee_pattern = next((p for p in patterns if p["direction"] == "出金" and p["account"] == FEE_ACCOUNT), None)
    fee_dates = {r["date"] for r in targets if r["withdrawal"] and r["withdrawal"] <= FEE_MAX and is_fee_text(r["text"])}
    lines = []
    for row, (result, _) in sorted(zip(targets, outcomes), key=lambda t: t[0]["no"]):  # 借入の履歴を日付順に積むため
        lines.extend(make_lines(row, result, args.bank_account, args.bank_sub, fee_pattern, fee_dates))

    counts = {s: sum(1 for l in lines if l["status"] == s) for s in STATUS_ORDER}
    mismatches = [r for r in passbook if r["check"].startswith("不一致")]
    memos = [r for r in passbook if r["memo"]]
    summary = [
        ["■ 概要", ""],
        ["銀行口座", bank_label],
        ["通帳の取引数", len(targets)],
        ["仕訳行数", len(lines)],
        ["JEV", jev_note],
        ["判定にかかった時間", f"{elapsed:.1f}秒"],
        ["", ""],
        ["■ 判定の内訳", ""],
        *[[s, counts[s]] for s in STATUS_ORDER],
        ["", ""],
        ["■ 残高チェック", "OK" if not mismatches else f"{len(mismatches)}か所で不一致（通帳の読み取りミスの可能性）"],
        *[[f"通帳{r['no']}行目", r["check"]] for r in mismatches],
        *([["■ 読み取りに自信がない箇所", ""]] + [[f"通帳{r['no']}行目", r["memo"]] for r in memos] if memos else []),
        ["", ""],
        ["■ 判定の意味", ""],
        ["過去一致", "過去の元帳に同じ相手先があり、科目が1つに決まる"],
        ["JEV判定", f"JEVが候補から選び、確信度が{args.threshold:.0%}以上"],
        ["推定", "科目は分かるが、元本と利息・手数料の内訳を推定した"],
        ["要確認", "候補はあるが自信が低い。必ず確認する"],
        [UNKNOWN, "判断できない。科目は「不明」のまま"],
    ]
    write_excel(args.output, lines, passbook, rules, summary)

    print(f"仕訳 {len(lines)}行を出力しました: {args.output}")
    print("  " + " / ".join(f"{s} {counts[s]}" for s in STATUS_ORDER))
    print(f"  JEV: {jev_note}（{elapsed:.1f}秒）")
    print(f"  残高チェック: {'OK' if not mismatches else f'{len(mismatches)}か所で不一致'}")
    for l in lines:
        if l["status"] in ("要確認", UNKNOWN):
            print(f"  [{l['status']}] 通帳{l['no']}行目 {l['date']} {l['text']} {l['amount']:,}円 … {l['reason']}")
    if errors:
        print("JEVのエラー: " + errors[0], file=sys.stderr)


if __name__ == "__main__":
    main()
